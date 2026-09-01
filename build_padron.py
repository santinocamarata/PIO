"""Genera padron.js a partir del Excel de alumnos a ingresar.

Uso:
    python build_padron.py "ruta/al/Alumnos_a_ingresar_XXXXXXX.xlsx"

Salida: padron.js con `window.PADRON = {"<legajo>": "<Apellido, Nombre>", ...}`
Lookup O(1) por legajo normalizado.

El nombre NO se parte en nombre/apellido: el Excel de origen trae un solo campo
y alrededor del 8% de las filas vienen con el orden invertido, asi que cualquier
split automatico corrompe esos registros.

Este script es opcional. La aplicacion tiene un boton "Actualizar padron" que
lee el mismo Excel directamente desde el navegador. Usar este script solo para
regenerar el padron embebido que viene por defecto.

Requiere: pip install openpyxl
"""
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

import openpyxl

LARGO_LEGAJO = 7
DESTINO = Path(__file__).with_name("padron.js")


def normalizar_legajo(valor):
    """Deja solo digitos. 1251113 / '1251113' / ' 1251113 ' -> '1251113'."""
    return re.sub(r"\D", "", str(valor))


def normalizar_nombre(valor):
    """Trim, colapsa espacios repetidos y limpia el espacio previo a la coma."""
    texto = unicodedata.normalize("NFC", str(valor))
    texto = re.sub(r"\s+", " ", texto).strip()
    texto = re.sub(r"\s+,", ",", texto)
    texto = re.sub(r",\s*", ", ", texto)
    return texto


def ubicar_columnas(encabezado):
    """Devuelve (indice_legajo, indice_alumno, fila_inicial)."""
    celdas = [str(c or "").strip().lower() for c in encabezado]
    i_legajo = next((i for i, c in enumerate(celdas) if "legajo" in c or c == "lu"), None)
    i_alumno = next(
        (i for i, c in enumerate(celdas) if "alumno" in c or "nombre" in c or "apellido" in c),
        None,
    )
    if i_legajo is None or i_alumno is None:
        return 0, 1, 0
    return i_legajo, i_alumno, 1


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1

    origen = Path(sys.argv[1])
    if not origen.is_file():
        print(f"No se encontro el archivo: {origen}")
        return 1

    libro = openpyxl.load_workbook(origen, read_only=True, data_only=True)
    hoja = libro[libro.sheetnames[0]]
    filas = list(hoja.iter_rows(values_only=True))
    libro.close()

    if not filas:
        print("El archivo no tiene filas.")
        return 1

    i_legajo, i_alumno, desde = ubicar_columnas(filas[0])

    padron = {}
    descartadas = []
    for numero, fila in enumerate(filas[desde:], start=desde + 1):
        legajo = normalizar_legajo(fila[i_legajo]) if i_legajo < len(fila) else ""
        nombre = normalizar_nombre(fila[i_alumno]) if i_alumno < len(fila) else ""
        if len(legajo) != LARGO_LEGAJO or not nombre:
            descartadas.append(numero)
            continue
        padron[legajo] = nombre

    if not padron:
        print("No se encontro ningun legajo valido. Revisa las columnas del archivo.")
        return 1

    cuerpo = json.dumps(padron, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
    meta = {"origen": origen.name, "cantidad": len(padron), "generado": date.today().isoformat()}

    DESTINO.write_text(
        "/* Padron de alumnos a ingresar. Archivo generado: no editar a mano.\n"
        '   Regenerar con: python build_padron.py "<archivo.xlsx>" */\n'
        f"window.PADRON_META = {json.dumps(meta, ensure_ascii=False)};\n"
        f"window.PADRON = {cuerpo};\n",
        encoding="utf-8",
    )

    print(f"escrito: {DESTINO}")
    print(f"registros: {len(padron)}   descartados: {len(descartadas)} {descartadas[:5]}")
    print(f"peso: {DESTINO.stat().st_size / 1024:.1f} KB")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
